"""Format-specific readers.

Each reader's only job is to turn a file into (row_number, record_dict) pairs
and land them verbatim in stg_raw_record. No cleaning, no type coercion, no
mapping - those belong in harmonise.py, where they can be changed without
re-reading source files.

Adding a new file *format* means adding a reader here. Adding a new *country*
in an existing format means adding an entry to config/sources.yml only - by
hand, or through the Ingest tab, which uses the peek/guess helpers at the end
of this module to propose the entry from the file itself.
"""
from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook
from rapidfuzz import fuzz

from . import db, harmonise

CANONICAL_FIELDS = [
    "source_record_id", "txn_date", "ministry_code", "ministry_name", "account_code",
    "description", "supplier", "amount", "currency", "payment_method", "fiscal_year",
]


# --------------------------------------------------------------- readers ------

def default_data_dir() -> Path:
    """Find candidate_data inside the repo or up to two levels beside it.

    Both layouts work without --data-dir. Falls back to the in-repo path so
    the caller can report where it looked.
    """
    for candidate in (db.ROOT / "candidate_data",
                      db.ROOT.parent / "candidate_data",
                      db.ROOT.parent.parent / "candidate_data"):
        if candidate.is_dir():
            return candidate
    return db.ROOT / "candidate_data"      # does not exist; caller reports it


def source_path(spec: dict, data_dir: Path) -> Path:
    """Where a country's file lives.

    Supplied extracts sit in the data directory given on the command line. A
    country added through the Ingest tab keeps its upload under the repo
    (`data_dir` in its config block), so the same entry works from both the
    app and `run_pipeline.py`.
    """
    base = (db.ROOT / spec["data_dir"]) if spec.get("data_dir") else data_dir
    return base / spec["path"]


def read_csv_source(path: Path, spec: dict) -> Iterator[tuple[int, dict]]:
    """Delimited text. Row numbers are file line numbers, header on line 1."""
    with path.open(encoding=spec.get("encoding", "utf-8-sig"), newline="") as fh:
        # Fields beyond the header are kept under "_extra" rather than a None
        # key, so a row with a stray delimiter still stages as valid JSON.
        reader = csv.DictReader(fh, delimiter=spec.get("delimiter") or ",", restkey="_extra")
        for i, rec in enumerate(reader, start=2):
            yield i, rec


def read_excel_source(path: Path, spec: dict) -> Iterator[tuple[int, dict]]:
    """Excel reader that copes with a report banner above the real header.

    Country B's sheet carries five lines of SIGFiP letterhead, a blank line,
    then the header on row 7, and closes with a printed TOTAL line. The header
    row is taken from config but verified: if the configured row does not look
    like a header, the sheet is scanned for one, so a small change in the number
    of banner lines next month does not silently shift every column.
    """
    ws = load_workbook(path, data_only=True)[spec["sheet"]]
    header_row = int(spec.get("header_row", 1))
    expected = {str(v).strip().lower() for v in spec["field_map"].values()}

    def header_at(r: int) -> list[str] | None:
        # A row is the header if at least half the mapped column names (min 2) sit in it.
        vals = next(ws.iter_rows(min_row=r, max_row=r, values_only=True), None)
        if not vals:
            return None
        cells = [str(v).strip() if v is not None else "" for v in vals]
        if len({c.lower() for c in cells} & expected) >= max(2, len(expected) // 2):
            return cells
        return None

    header = header_at(header_row)
    if header is None:
        for r in range(1, min(ws.max_row, 30) + 1):
            header = header_at(r)
            if header:
                header_row = r
                break
    if header is None:
        raise ValueError(f"{path.name}: could not locate a header row in '{spec['sheet']}'")

    marker = str(spec.get("total_row_marker") or "").strip().upper()
    for r in range(header_row + 1, ws.max_row + 1):
        vals = next(ws.iter_rows(min_row=r, max_row=r, values_only=True), None)
        if vals is None:
            continue
        if all(v is None or str(v).strip() == "" for v in vals):
            continue
        rec = {h: v for h, v in zip(header, vals) if h}
        first = str(vals[0]).strip().upper() if vals[0] is not None else ""
        if marker and first == marker:
            # printed control total - surfaced to the batch, never a fact row
            rec["__is_total_row__"] = True
        yield r, rec


def read_json_source(path: Path, spec: dict) -> Iterator[tuple[int, dict]]:
    """A JSON array of objects, at the top level or under a dotted `records_path`
    ("data.records"). Row numbers are 1-based positions in the array. Nested
    sub-transactions stay inside the record; harmonise splits them out."""
    doc = json.loads(path.read_text(encoding="utf-8"))
    records = doc
    for key in str(spec.get("records_path", "")).split("."):
        if key:
            records = records[key]
    for i, rec in enumerate(records, start=1):
        yield i, rec


READERS = {"csv": read_csv_source, "excel": read_excel_source, "json": read_json_source}
READER_FOR_SUFFIX = {".csv": "csv", ".txt": "csv", ".xlsx": "excel", ".xlsm": "excel",
                     ".json": "json"}


def read_source_metadata(path: Path, spec: dict) -> dict | None:
    """The source's own metadata block, if the config names one. Country C
    declares a record count, a currency and an extraction timestamp there;
    the count is reconciled after staging and the timestamp bounds the dates."""
    if spec.get("reader") != "json" or not spec.get("metadata_path"):
        return None
    doc = json.loads(path.read_text(encoding="utf-8"))
    return doc.get(spec["metadata_path"])


def read_country_chart_of_accounts(path: Path, spec: dict) -> list[tuple[str, str]]:
    """(code, label) pairs from a chart-of-accounts sheet in the same workbook,
    if the config names one (Country B's Plan_comptable). Column names come
    from the config; if they are not found, the first two columns are used."""
    if not spec.get("chart_of_accounts_sheet"):
        return []
    ws = load_workbook(path, data_only=True)[spec["chart_of_accounts_sheet"]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(v).strip() if v is not None else "" for v in rows[0]]
    try:
        ci = header.index(spec["chart_of_accounts_code_column"])
        li = header.index(spec["chart_of_accounts_label_column"])
    except ValueError:
        ci, li = 0, 1
    out = []
    for r in rows[1:]:
        if r and r[ci] is not None:
            out.append((str(r[ci]).strip(), str(r[li]).strip() if r[li] is not None else None))
    return out


# ----------------------------------------------------------------- staging ----

def ingest_country(conn, country_code: str, spec: dict, data_dir: Path) -> int:
    """Stage one country's file verbatim and return the batch id.

    Opens a batch, loads the chart of accounts if there is one, then writes
    every row's payload as JSON with its file/sheet/row position. A printed
    TOTAL row becomes the batch's control total instead of a staged row. The
    row count is checked against the source's own declared count where it
    gives one.
    """
    path = source_path(spec, data_dir)
    if not path.exists():
        raise FileNotFoundError(path)

    metadata = read_source_metadata(path, spec)
    cur = conn.execute(
        "INSERT INTO ingestion_batch (country_code, source_file, source_system, started_at,"
        " source_metadata) VALUES (?,?,?,?,?)",
        (country_code, spec["path"], spec.get("source_system"), db.utc_now(),
         json.dumps(metadata) if metadata else None),
    )
    batch_id = cur.lastrowid

    # A country-supplied chart of accounts is authoritative for account labels.
    chart_pairs = read_country_chart_of_accounts(path, spec)
    if chart_pairs:
        db.load_country_chart_of_accounts(conn, country_code, chart_pairs)

    reader = READERS[spec["reader"]]
    id_field = spec["field_map"].get("source_record_id")
    rows_read = 0
    rows_loaded = 0
    control_total = None
    sheet = spec.get("sheet")

    for row_no, rec in reader(path, spec):
        rows_read += 1

        # printed total line: capture as a control figure, do not stage as data
        if rec.pop("__is_total_row__", False):
            amt_field = spec["field_map"].get("amount")
            raw = rec.get(amt_field)
            # Same parser as the transactions, so a printed "49 626 000 570,00"
            # is read under the same conventions as the rows it totals.
            control_total, _ = harmonise.parse_amount(raw, spec.get("conventions", {}) or {})
            db.add_issue(
                conn, rule_code="SRC_CONTROL_TOTAL", severity="INFO",
                country_code=country_code, batch_id=batch_id, field=amt_field,
                raw_value=raw,
                message=f"Source printed a control total of {control_total!r}; "
                        f"excluded from transactions and reconciled after load.",
            )
            continue

        src_id = str(rec.get(id_field)).strip() if rec.get(id_field) is not None else None
        conn.execute(
            "INSERT INTO stg_raw_record (batch_id, country_code, source_file, source_sheet,"
            " source_row_no, source_record_id, payload_json, ingested_at) VALUES (?,?,?,?,?,?,?,?)",
            (batch_id, country_code, spec["path"], sheet, row_no, src_id,
             json.dumps(rec, default=str, ensure_ascii=False), db.utc_now()),
        )
        rows_loaded += 1

    conn.execute(
        "UPDATE ingestion_batch SET rows_read=?, rows_loaded=?, control_total=? WHERE batch_id=?",
        (rows_read, rows_loaded, control_total, batch_id),
    )

    # Reconcile the source's own declared record count where it gives one.
    if metadata and metadata.get("recordCount") is not None:
        declared = int(metadata["recordCount"])
        if declared != rows_loaded:
            db.add_issue(
                conn, rule_code="SRC_COUNT_MISMATCH", severity="WARN",
                country_code=country_code, batch_id=batch_id,
                message=f"Source metadata declares {declared} records; {rows_loaded} staged.",
            )
    conn.commit()
    return batch_id


# ------------------------------------------------- new-country helpers --------
# Used by the Ingest tab to propose a sources.yml entry from the file itself.
# Everything here is a *suggestion* the analyst confirms; nothing is loaded
# from a guess.

def peek_source(path: Path, reader: str, options: dict | None = None,
                sample: int = 5) -> dict:
    """Column names and a few sample rows, without staging anything.

    Returns {"columns": [...], "sample": [dict, ...], "sheets": [...],
             "nested": [keys whose values are lists of dicts]}.
    """
    options = options or {}
    out: dict = {"columns": [], "sample": [], "sheets": [], "nested": []}

    if reader == "csv":
        with path.open(encoding=options.get("encoding", "utf-8-sig"), newline="") as fh:
            rd = csv.DictReader(fh, delimiter=options.get("delimiter") or guess_delimiter(path),
                                restkey="_extra")
            out["columns"] = [str(c) for c in (rd.fieldnames or [])]
            for i, rec in enumerate(rd):
                if i >= sample:
                    break
                out["sample"].append({str(k): v for k, v in rec.items()})

    elif reader == "excel":
        wb = load_workbook(path, read_only=True, data_only=True)
        out["sheets"] = wb.sheetnames
        ws = wb[options.get("sheet") or wb.sheetnames[0]]
        header_row = int(options.get("header_row") or 1)
        rows = ws.iter_rows(min_row=header_row, max_row=header_row + sample, values_only=True)
        header = next(rows, None) or ()
        out["columns"] = [str(v).strip() for v in header if v is not None and str(v).strip()]
        for vals in rows:
            if vals and any(v is not None for v in vals):
                out["sample"].append({h: v for h, v in zip(header, vals) if h})
        wb.close()

    elif reader == "json":
        doc = json.loads(path.read_text(encoding="utf-8"))
        records = doc
        for key in str(options.get("records_path") or "").split("."):
            if key:
                records = records[key]
        if isinstance(records, dict):
            # Not a list: offer the keys that hold lists as records_path candidates.
            out["nested"] = [k for k, v in records.items() if isinstance(v, list)]
            records = []
        first = records[0] if records else {}
        out["columns"] = [k for k, v in first.items() if not isinstance(v, (list, dict))]
        out["nested"] = out["nested"] or [
            k for k, v in first.items() if isinstance(v, list) and v and isinstance(v[0], dict)]
        for rec in records[:sample]:
            out["sample"].append({k: v for k, v in rec.items() if not isinstance(v, (list, dict))})

    return out


# What each canonical field tends to be called. English and French because
# those are the languages in hand; extend as countries are added.
FIELD_HINTS = {
    "source_record_id": ["txn_id", "transaction_id", "id_transaction", "transactionid",
                         "id", "ref", "reference", "voucher", "piece"],
    "txn_date": ["date", "txn_date", "posting_date", "postingdate", "date_ecriture",
                 "date_paiement", "value_date"],
    "ministry_code": ["ministry_code", "ministere_code", "ministrycode", "mda_code",
                      "vote_code", "vote"],
    "ministry_name": ["ministry_name", "ministere_nom", "ministryname", "ministry",
                      "ministere", "mda", "vote_name"],
    "account_code": ["account_code", "code_budgetaire", "coacode", "account", "gl_code",
                     "gl_account", "chart_code", "economic_code", "compte", "item_code"],
    "description": ["description", "libelle", "narration", "memo", "details", "objet",
                    "purpose", "particulars"],
    "supplier": ["vendor", "supplier", "tiers", "payee", "fournisseur", "beneficiaire"],
    "amount": ["amount", "montant", "value", "amt", "total", "net_amount"],
    "currency": ["currency", "devise", "ccy", "monnaie"],
    "payment_method": ["payment_method", "mode_paiement", "pay_mode", "payment_mode",
                       "payment_type"],
    "fiscal_year": ["fiscal_year", "fiscalyear", "exercice", "fy", "annee_budgetaire"],
}


def guess_field_map(columns: list[str], threshold: int = 75) -> dict[str, str]:
    """Propose canonical_field -> source column from column names alone.

    A column is offered to at most one field, the best-scoring one first, so
    "MINISTRY_CODE" cannot be claimed by both ministry_code and ministry_name.
    """
    scored: list[tuple[float, str, str]] = []
    for col in columns:
        folded = harmonise.fold(col).replace(" ", "_")
        for field, hints in FIELD_HINTS.items():
            best = 0.0
            for h in hints:
                if folded == h:
                    best = 100.0
                    break
                if folded.startswith(h + "_") or folded.endswith("_" + h):
                    best = max(best, 95.0)          # amount_kes, id_transaction
                else:
                    best = max(best, fuzz.ratio(folded, h))
            if best >= threshold:
                scored.append((best, field, col))
    scored.sort(reverse=True)
    out: dict[str, str] = {}
    used: set[str] = set()
    for _, field, col in scored:
        if field not in out and col not in used:
            out[field] = col
            used.add(col)
    return out


def guess_delimiter(path: Path, encoding: str = "utf-8-sig") -> str:
    """The delimiter that splits the header line into the most fields."""
    with path.open(encoding=encoding, newline="") as fh:
        header = fh.readline()
    counts = {d: header.count(d) for d in (",", ";", "\t", "|")}
    best = max(counts, key=counts.get)
    return best if counts[best] else ","


DATE_SHAPES = [
    (re.compile(r"^\d{4}-\d{2}-\d{2}"), "%Y-%m-%d"),
    (re.compile(r"^\d{2}/\d{2}/\d{4}$"), "%d/%m/%Y"),
    (re.compile(r"^\d{2}-\d{2}-\d{4}$"), "%d-%m-%Y"),
    (re.compile(r"^\d{2}\.\d{2}\.\d{4}$"), "%d.%m.%Y"),
    (re.compile(r"^\d{4}/\d{2}/\d{2}$"), "%Y/%m/%d"),
    (re.compile(r"^\d{8}$"), "%Y%m%d"),
]


def guess_date_format(values: list) -> str | None:
    """The strptime format most sample values fit, or None. Day-first is
    assumed for dd/mm/yyyy shapes - the analyst confirms."""
    votes: dict[str, int] = {}
    for v in values:
        if v is None:
            continue
        s = str(v).strip()
        for rx, fmt in DATE_SHAPES:
            if rx.match(s):
                votes[fmt] = votes.get(fmt, 0) + 1
                break
    return max(votes, key=votes.get) if votes else None


def guess_conventions(values: list) -> dict:
    """Decimal/thousands conventions suggested by the sample amounts."""
    strs = [str(v) for v in values if v is not None and not isinstance(v, (int, float))]
    comma_dec = sum(1 for s in strs if re.search(r",\d{1,2}$", s))
    dot_dec = sum(1 for s in strs if re.search(r"\.\d{1,2}$", s))
    suffixes = sorted({m.group(1) for s in strs for m in [re.search(r"\s([A-Z]{2,5})$", s)] if m})
    return {
        "decimal_separator": "," if comma_dec > dot_dec else ".",
        "thousands_separators": [" ", "."] if comma_dec > dot_dec else [","],
        "strip_suffixes": suffixes,
    }
