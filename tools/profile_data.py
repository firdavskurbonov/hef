"""One-off exploration of the three supplied extracts.

Not part of the prototype. Answers "what is actually in these files?" before
any design decision; the findings are written up in DATA_QUALITY_FINDINGS.md.
Prints to stdout, changes nothing.
"""
import json
import os
import re
from collections import Counter

import pandas as pd
from openpyxl import load_workbook

def _find_data() -> str:
    """candidate_data inside the repo or up to two levels beside it."""
    here = os.path.dirname(os.path.abspath(__file__))
    for rel in ("candidate_data",
                os.path.join("..", "candidate_data"),
                os.path.join("..", "..", "candidate_data")):
        p = os.path.normpath(os.path.join(here, rel))
        if os.path.isdir(p):
            return p
    raise SystemExit(
        "candidate_data/ not found. Place the supplied extracts in "
        "candidate_data/ at the repository root (or beside the repo) and re-run."
    )


DATA = _find_data()
pd.set_option("display.width", 200)
pd.set_option("display.max_columns", 50)


def rule(t):
    """Section banner."""
    print("\n" + "=" * 78)
    print(t)
    print("=" * 78)


def col_profile(df, name):
    """Per column: nulls, blanks, distinct count and three sample values."""
    print(f"\nrows={len(df)}  cols={len(df.columns)}   [{name}]")
    for c in df.columns:
        s = df[c]
        n_null = int(s.isna().sum())
        n_blank = int((s.astype(str).str.strip() == "").sum())
        uniq = s.nunique(dropna=True)
        samples = [repr(v) for v in s.dropna().unique()[:3]]
        print(
            f"  {c:<22} null={n_null:<5} blank={n_blank:<5} uniq={uniq:<6} "
            f"e.g. {', '.join(samples)}"
        )


def text_anomalies(series, name, max_distinct=60):
    """Cardinality, frequency and length of a text column.

    A free-text column that behaves like a closed list is worth listing in
    full: the values that appear once, or that are far longer than the rest,
    are where anything unusual hides. Needs no knowledge of what the text says,
    which is the point - it is the only check that catches text written to
    mislead a human reader rather than a parser.
    """
    vals = series.dropna().astype(str)
    counts = vals.value_counts()
    print(f"\n[{name}]  {len(vals)} rows, {len(counts)} distinct")
    if len(counts) > max_distinct:
        print(f"  high cardinality - genuine free text; length "
              f"median={int(vals.str.len().median())} max={int(vals.str.len().max())}")
        return
    lengths = counts.index.to_series().str.len()
    typical = lengths.median()
    print(f"  closed list - length median={int(typical)} max={int(lengths.max())}")
    for value, n in counts.sort_values().items():
        flag = "  <== OUTLIER" if len(value) > 2 * typical else ""
        print(f"  {n:>5} x  len={len(value):>4}  {value[:88]!r}{flag}")


# ----------------------------------------------------------------- Country A
rule("COUNTRY A - country_a_expenditure.csv")
a = pd.read_csv(os.path.join(DATA, "country_a_expenditure.csv"), dtype=str)
col_profile(a, "A raw (all str)")

print("\n-- DATE parsing (declared dd/mm/yyyy) --")
d_dmy = pd.to_datetime(a["DATE"], format="%d/%m/%Y", errors="coerce")
print(f"  parsed as dd/mm/yyyy : {d_dmy.notna().sum()} ok, {d_dmy.isna().sum()} failed")
bad = a.loc[d_dmy.isna(), "DATE"].dropna().unique()[:15]
print(f"  unparseable samples  : {list(bad)}")
d_any = pd.to_datetime(a["DATE"], errors="coerce", dayfirst=True)
print(f"  loose parse          : {d_any.notna().sum()} ok")
print(f"  date range           : {d_any.min()} -> {d_any.max()}")

print("\n-- AMOUNT_KES --")
amt = pd.to_numeric(a["AMOUNT_KES"], errors="coerce")
print(f"  numeric ok={amt.notna().sum()}  non-numeric={amt.isna().sum()}")
print(f"  non-numeric samples: {list(a.loc[amt.isna(),'AMOUNT_KES'].dropna().unique()[:10])}")
print(f"  negatives={int((amt < 0).sum())}  zeros={int((amt == 0).sum())}")
print(f"  min={amt.min()}  max={amt.max()}  sum={amt.sum():,.2f}")

print("\n-- keys / codes --")
print(f"  TXN_ID duplicated rows: {int(a['TXN_ID'].duplicated().sum())}")
dups = a[a["TXN_ID"].duplicated(keep=False)].sort_values("TXN_ID")
if len(dups):
    print(dups.head(6).to_string(index=False))
print(f"  ACCOUNT_CODE lengths  : {Counter(a['ACCOUNT_CODE'].astype(str).str.len())}")
print(f"  MINISTRY_CODE values  : {dict(Counter(a['MINISTRY_CODE'].dropna()))}")
print(f"  PAYMENT_METHOD values : {dict(Counter(a['PAYMENT_METHOD'].dropna()))}")
print(f"  DESCRIPTION blank/na  : {int(a['DESCRIPTION'].isna().sum())}")

rule("COUNTRY A - free-text anomaly scan")
text_anomalies(a["DESCRIPTION"], "A DESCRIPTION")
text_anomalies(a["VENDOR"], "A VENDOR")

# ----------------------------------------------------------------- Country B
rule("COUNTRY B - country_b_depenses.xlsx")
wb = load_workbook(os.path.join(DATA, "country_b_depenses.xlsx"), data_only=True)
print(f"sheets: {wb.sheetnames}")
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"  '{sn}': dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
    for row in ws.iter_rows(min_row=1, max_row=4, values_only=True):
        print("     ", row)

b = pd.read_excel(os.path.join(DATA, "country_b_depenses.xlsx"), sheet_name="Depenses", dtype=str)
col_profile(b, "B Depenses raw")

# The check that surfaced the injected instructions. libelle is effectively the
# chart-of-accounts label, so it behaves as a closed list - and four values
# appear exactly once, several times longer than any other.
rule("COUNTRY B - free-text anomaly scan")
_bh = pd.read_excel(os.path.join(DATA, "country_b_depenses.xlsx"),
                    sheet_name="Depenses", dtype=str, header=6)
_bh = _bh[_bh.iloc[:, 0].astype(str).str.upper() != "TOTAL"]      # drop the printed total
for _c in _bh.columns:
    if re.search(r"libell|descr", str(_c), re.I):
        text_anomalies(_bh[_c], f"B {_c}")
chart_b = pd.read_excel(os.path.join(DATA, "country_b_depenses.xlsx"), sheet_name="Plan_comptable", dtype=str)
col_profile(chart_b, "B Plan_comptable")
print("\nPlan_comptable full contents:")
print(chart_b.to_string(index=False))

for c in b.columns:
    if re.search(r"date", c, re.I):
        print(f"\n-- B date column '{c}' --")
        print(f"  samples: {list(b[c].dropna().unique()[:10])}")
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y"):
            ok = pd.to_datetime(b[c], format=fmt, errors="coerce").notna().sum()
            print(f"    {fmt}: {ok}")
    if re.search(r"montant|amount|valeur", c, re.I):
        print(f"\n-- B amount column '{c}' --")
        raw = b[c].dropna().astype(str)
        print(f"  samples: {list(raw.unique()[:10])}")
        num = pd.to_numeric(b[c], errors="coerce")
        print(f"  numeric ok={num.notna().sum()} bad={num.isna().sum()}")
        print(f"  bad samples: {list(b.loc[num.isna() & b[c].notna(), c].unique()[:10])}")

# ----------------------------------------------------------------- Country C
rule("COUNTRY C - country_c_expenditure.json")
with open(os.path.join(DATA, "country_c_expenditure.json"), encoding="utf-8") as fh:
    cj = json.load(fh)
print("metadata:")
for k, v in cj["metadata"].items():
    print(f"  {k}: {v}")
tx = cj["transactions"]
print(f"\ntransactions: {len(tx)} (metadata claims {cj['metadata'].get('recordCount')})")

keys = Counter()
for t in tx:
    keys.update(t.keys())
print("\nkey frequency across records:")
for k, v in keys.most_common():
    print(f"  {k:<24} {v}  ({'always' if v == len(tx) else 'SOMETIMES'})")

nested = [t for t in tx if isinstance(t.get("subTransactions"), list) and t["subTransactions"]]
print(f"\nrecords with non-empty subTransactions: {len(nested)}")
if nested:
    print("example parent + children:")
    print(json.dumps(nested[0], indent=2)[:1200])
    sub_keys = Counter()
    n_sub = 0
    for t in nested:
        for st in t["subTransactions"]:
            sub_keys.update(st.keys())
            n_sub += 1
    print(f"\n  total sub-transactions: {n_sub}")
    print(f"  sub keys: {dict(sub_keys)}")
    # do children sum to the parent?
    mism = 0
    for t in nested:
        try:
            tot = sum(float(st.get("amount", 0) or 0) for st in t["subTransactions"])
            p = float(t.get("amount", 0) or 0)
            if abs(tot - p) > 0.01:
                mism += 1
        except (TypeError, ValueError):
            mism += 1
    print(f"  parents where children do NOT sum to parent amount: {mism} / {len(nested)}")

cf = pd.json_normalize(tx)
col_profile(cf.drop(columns=[c for c in cf.columns if c == "subTransactions"], errors="ignore"), "C flattened")

rule("COUNTRY C - free-text anomaly scan")
if "description" in cf:
    text_anomalies(cf["description"], "C description")

for c in cf.columns:
    if c == "subTransactions":
        continue
    if re.search(r"currency|amount|rate", c, re.I):
        vals = cf[c].dropna()
        print(f"\n-- C '{c}' --  uniq={vals.nunique()}")
        print(f"   samples: {list(vals.unique()[:12])}")
print("\n-- C dates --")
for c in [x for x in cf.columns if re.search(r"date", x, re.I)]:
    print(f"  {c}: samples={list(cf[c].dropna().unique()[:8])}")
    print(f"     iso-parse ok={pd.to_datetime(cf[c], errors='coerce', format='ISO8601').notna().sum()}")
print(f"\n  transactionId duplicated: {int(cf['transactionId'].duplicated().sum())}")

# ------------------------------------------------------------- cross-cutting
rule("CROSS-CUTTING")
print("Ministry codes/names per country (for a ministry dimension):")
print("  A:", dict(Counter(a["MINISTRY_NAME"].dropna())))
mcol = [c for c in b.columns if re.search(r"minist", c, re.I)]
if mcol:
    print("  B:", dict(Counter(b[mcol[0]].dropna())))
print("  C:", dict(Counter(cf["ministryName"].dropna())) if "ministryName" in cf else "n/a")

print("\nDescription vocabulary overlap (rough classification signal):")
def toks(series):
    """Word counts over a text column, lowercased, words over three letters."""
    out = Counter()
    for v in series.dropna().astype(str):
        out.update(w for w in re.findall(r"[a-zA-Zàâçéèêëîïôûùüÿœ]+", v.lower()) if len(w) > 3)
    return out
ta = toks(a["DESCRIPTION"])
dcol_b = [c for c in b.columns if re.search(r"libell|descr", c, re.I)]
tb = toks(b[dcol_b[0]]) if dcol_b else Counter()
tc = toks(cf["description"]) if "description" in cf else Counter()
print(f"  A top: {[w for w, _ in ta.most_common(18)]}")
print(f"  B top: {[w for w, _ in tb.most_common(18)]}")
print(f"  C top: {[w for w, _ in tc.most_common(18)]}")
